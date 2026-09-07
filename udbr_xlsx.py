#!/usr/bin/env python3
"""
udbr_xlsx — the house style for every workbook this pipeline produces.

The fonts, fills, borders, header block and table writer were duplicated
between build_workbook.py and build_design_workbook.py. They were identical
except where they had drifted: one wrapped long text and the other did not, and
the two spelled the same severity colours in two places.

A workbook that goes to the RCM team should not look different depending on
which script built it, and a colour change should not be two edits.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = 'Arial'

H1    = Font(name=FONT, size=14, bold=True, color='1F3864')
H2    = Font(name=FONT, size=11, bold=True, color='FFFFFF')
BODY  = Font(name=FONT, size=10)
BOLD  = Font(name=FONT, size=10, bold=True)
NOTE  = Font(name=FONT, size=9, italic=True, color='595959')

HFILL = PatternFill('solid', fgColor='1F3864')

# Severity carries colour AND the word, never colour alone: a workbook printed
# in greyscale or read by someone colour-blind must still be sortable.
SEV = {'High':   PatternFill('solid', fgColor='F8CBAD'),
       'Medium': PatternFill('solid', fgColor='FFE699'),
       'Low':    PatternFill('solid', fgColor='E2EFDA')}

THIN = Border(*[Side(style='thin', color='BFBFBF')] * 4)


def header(ws, title, subtitle, provenance=None):
    """Title block. `provenance` is the line naming the source and date — every
    workbook needs one, because a findings list without its as-of date is
    indistinguishable from a current one."""
    ws['A1'] = title;    ws['A1'].font = H1
    ws['A2'] = subtitle; ws['A2'].font = NOTE
    if provenance:
        ws['A3'] = provenance; ws['A3'].font = NOTE


def write_table(ws, dfr, start, name, widths, wrap=()):
    """A bordered, filterable table with a styled header row.

    `wrap` names the columns whose text should wrap — findings text and
    evidence, generally. The two builders differed on this, so the same finding
    rendered readably in one workbook and as a single clipped line in the
    other.

    Returns the header row index."""
    cols = list(dfr.columns)
    for j, col in enumerate(cols, 1):
        c = ws.cell(start, j, col)
        c.font = H2
        c.fill = HFILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = THIN

    for i, (_, row) in enumerate(dfr.iterrows(), start + 1):
        for j, col in enumerate(cols, 1):
            c = ws.cell(i, j, row[col])
            c.font = BODY
            c.border = THIN
            c.alignment = Alignment(vertical='top', wrap_text=(col in wrap))
        if 'Severity' in cols and row['Severity'] in SEV:
            ws.cell(i, cols.index('Severity') + 1).fill = SEV[row['Severity']]

    if len(dfr):
        t = Table(displayName=name,
                  ref=f'A{start}:{get_column_letter(len(cols))}{start + len(dfr)}')
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
        ws.add_table(t)

    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(start + 1, 1)
    # The header row, so a caller can add a column beside the table without
    # recomputing where it starts.
    return start


def finish(wb):
    """Applied to every workbook before saving, so none of them ships with
    gridlines showing when the others do not."""
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 20
    return wb
