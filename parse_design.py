#!/usr/bin/env python3
"""
Parse UDBR_Profile_Design_v9.xlsx into a flat rule table.

This is a DESIGN workbook, not a production dump, and its shape is different:
one sheet per profile, rules grouped into bands by the AOH level the design
places them at, with a 'shared' flag marking rules common to both alphas.

The band a rule sits under IS the design decision. Nothing else in the sheet
records intended level, so the parse has to track which band it is inside.
"""
import udbr_paths as P
from udbr_domain import TIERS, TIER_RANK, tier_rank
import sys, json, re
from openpyxl import load_workbook

SRC = sys.argv[1] if len(sys.argv) > 1 else '/mnt/user-data/uploads/UDBR_Profile_Design_v9__1_.xlsx'
LEVELS = TIERS

wb = load_workbook(SRC, data_only=True)
rules, profiles = [], []

for name in wb.sheetnames:
    if name == 'Summary':
        continue
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    title = str(rows[0][0]).strip() if rows and rows[0][0] else name
    # 'Design Complete' flag lives on row 2, column D
    complete = ''
    for r in rows[:4]:
        for i, c in enumerate(r):
            if c and str(c).strip() == 'Design Complete':
                complete = str(r[i+1]).strip() if i+1 < len(r) and r[i+1] else ''
    # current-state block: tier / entity / rule count as the design found them
    cur = []
    hdr_i = None
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else '' for c in r]
        if cells[:3] == ['Tier', 'Entity', 'Lineage Path']:
            hdr_i = i
            j = i + 1
            while j < len(rows):
                c2 = [str(x).strip() if x is not None else '' for x in rows[j]]
                if c2[0] in LEVELS:
                    cur.append(dict(tier=c2[0], entity=c2[1], rules=c2[3],
                                    pid=c2[4] if len(c2) > 4 else ''))
                    j += 1
                else:
                    break
            break
    profiles.append(dict(sheet=name, name=title, complete=complete, current=cur))

    # rule table: find the column header row, then walk bands
    col_i = None
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else '' for c in r]
        if 'Target Field' in cells and 'Priority' in cells:
            col_i = i
            cols = cells
            break
    if col_i is None:
        continue
    ix = {c: k for k, c in enumerate(cols) if c}
    band = None; band_entity = None
    for r in rows[col_i+1:]:
        cells = [str(c).strip() if c is not None else '' for c in r]
        if not any(cells):
            continue
        first = cells[0]
        # a band header names a level and nothing else in the row is a rule
        m = re.match(r'^(System|Organization|Customer|Agency)\b', first)
        if m and not cells[ix.get('Target Field', 2)]:
            # Capture the ENTITY as well as the level. The band header reads
            # "System  -  AIM (System)" where one entity exists at that level,
            # but Customer and Agency bands in the Default Profile sheet are
            # BARE, because two entities exist at each and the design does not
            # separate them. Recording that as unattributable is the point:
            # grouping on level alone made two customers' rules look like a
            # conflict inside one profile.
            band = m.group(1)
            ent = first[m.end():].strip(' -\u2014\u2013')
            band_entity = ent if ent else None
            continue
        fld = cells[ix['Target Field']] if ix.get('Target Field') is not None else ''
        if not fld or fld == 'Target Field':
            continue
        g = lambda k: cells[ix[k]] if ix.get(k) is not None and ix[k] < len(cells) else ''
        rules.append(dict(
            sheet=name, profile=title, level=band, entity=band_entity,
            tab=g('Tab'), priority=g('Priority'), field=fld,
            otype=g('Outcome Type'), ovalue=g('Outcome Value'), raw=g('Raw Value'),
            cond=g('Conditions'), desc=g('Rule Description'),
            rid=g('Rule Id'), shared=(g('Shared').lower() == 'shared')))

json.dump(dict(source=SRC.split('/')[-1], profiles=profiles, rules=rules),
          open(P.DESIGN_JSON, 'w'), separators=(',', ':'))

print(f"profile sheets      : {len(profiles)}")
print(f"rules parsed        : {len(rules)}")
print(f"rules with no band  : {sum(1 for r in rules if not r['level'])}")
from collections import Counter
print("\nby designed level:")
for lv, n in Counter(r['level'] for r in rules).most_common():
    print(f"  {lv or '(none)':14} {n}")
print("\nshared flag:")
print(f"  shared     {sum(1 for r in rules if r['shared'])}")
print(f"  not shared {sum(1 for r in rules if not r['shared'])}")
print("\ndesign complete flag:")
for v, n in Counter(p['complete'] for p in profiles).most_common():
    print(f"  {v or '(blank)':10} {n}")
